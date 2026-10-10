"""Bytes to page text: PDF text layer, OCR fallback, DOCX, images."""

from __future__ import annotations

import io
import os
from typing import Any

SPARSE_CHARS = 40
GARBLE_LETTER_RATIO = 0.15
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
TEXT_EXTS = {".md", ".txt", ".csv", ".json"}
BINARY_EXTS = PDF_EXTS | DOCX_EXTS | IMAGE_EXTS


def ocr_always() -> bool:
    return (os.getenv("GRAPHRAG_OCR_ALWAYS") or "").strip().lower() in {"1", "true", "yes", "on"}


def parse_document(
    path: str | None = None,
    *,
    data: bytes | None = None,
    filename: str = "",
) -> dict[str, Any]:
    """Return pages, joined text, engine, and warnings. Never raises for missing OCR."""
    name = filename or (os.path.basename(path) if path else "upload.bin")
    ext = os.path.splitext(name)[1].lower()
    raw = data
    if raw is None:
        if not path:
            return _empty("no path or bytes", name)
        mode = "rb" if ext in BINARY_EXTS else "r"
        if mode == "rb":
            with open(path, "rb") as fh:
                raw = fh.read()
        else:
            with open(path, encoding="utf-8") as fh:
                text = fh.read()
            if ext == ".json":
                import json

                try:
                    text = json.dumps(json.loads(text), indent=2)
                except json.JSONDecodeError:
                    pass
            return {
                "filename": name,
                "pages": [{"page": 1, "text": text, "ocr": False}],
                "text": text,
                "engine": "utf8",
                "warnings": [],
            }

    if ext in PDF_EXTS:
        return _parse_pdf(raw, name)
    if ext in DOCX_EXTS:
        return _parse_docx(raw, name)
    if ext in IMAGE_EXTS:
        return _parse_image(raw, name)
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("utf-8", errors="replace")
    return {
        "filename": name,
        "pages": [{"page": 1, "text": text, "ocr": False}],
        "text": text,
        "engine": "utf8",
        "warnings": [],
    }


def _empty(reason: str, name: str) -> dict[str, Any]:
    return {
        "filename": name,
        "pages": [],
        "text": "",
        "engine": "none",
        "warnings": [reason],
    }


def _parse_pdf(raw: bytes, name: str) -> dict[str, Any]:
    warnings: list[str] = []
    pages: list[dict[str, Any]] = []
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        for i, page in enumerate(reader.pages, start=1):
            extracted = (page.extract_text() or "").strip()
            used_ocr = False
            if _needs_ocr(extracted):
                ocr_text, warn = _ocr_pdf_page(raw, i - 1)
                if warn:
                    warnings.append(f"page {i}: {warn}")
                if ocr_text.strip():
                    extracted = ocr_text.strip()
                    used_ocr = True
            pages.append({"page": i, "text": extracted, "ocr": used_ocr})
        engine = "pypdf+ocr" if any(p["ocr"] for p in pages) else "pypdf"
    except Exception as exc:  # noqa: BLE001
        warnings.append(f"pypdf failed: {exc}")
        return _empty("; ".join(warnings), name)
    text = "\n\n".join(p["text"] for p in pages if p["text"])
    return {
        "filename": name,
        "pages": pages,
        "text": text,
        "engine": engine,
        "warnings": warnings,
    }


def _parse_docx(raw: bytes, name: str) -> dict[str, Any]:
    warnings: list[str] = []
    try:
        from docx import Document

        doc = Document(io.BytesIO(raw))
        parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        text = "\n".join(parts)
    except Exception as exc:  # noqa: BLE001
        return _empty(f"docx failed: {exc}", name)
    ocr_bits: list[str] = []
    for blob in _docx_image_blobs(doc):
        img_text, warn = _ocr_image_bytes(blob)
        if warn:
            warnings.append(warn)
        if img_text.strip():
            ocr_bits.append(img_text.strip())
    if ocr_bits:
        extra = "\n".join(ocr_bits)
        text = f"{text}\n{extra}".strip() if text else extra
        engine = "python-docx+ocr"
        used_ocr = True
    else:
        engine = "python-docx"
        used_ocr = False
    return {
        "filename": name,
        "pages": [{"page": 1, "text": text, "ocr": used_ocr}],
        "text": text,
        "engine": engine,
        "warnings": warnings,
    }


def _docx_image_blobs(doc: Any) -> list[bytes]:
    blobs: list[bytes] = []
    try:
        for rel in doc.part.rels.values():
            if "image" in (getattr(rel, "reltype", "") or ""):
                part = getattr(rel, "target_part", None)
                blob = getattr(part, "blob", None) if part is not None else None
                if blob:
                    blobs.append(blob)
    except Exception:  # noqa: BLE001
        return blobs
    return blobs


def _parse_image(raw: bytes, name: str) -> dict[str, Any]:
    text, warn = _ocr_image_bytes(raw)
    warnings = [warn] if warn else []
    return {
        "filename": name,
        "pages": [{"page": 1, "text": text, "ocr": bool(text) or not warn}],
        "text": text,
        "engine": "ocr" if not warn else "ocr_skipped",
        "warnings": warnings,
    }


def _sparse(text: str) -> bool:
    alnum = sum(1 for c in (text or "") if c.isalnum())
    return alnum < SPARSE_CHARS


def _garbled(text: str) -> bool:
    """True when a text layer exists but is mostly not letters (mojibake / junk encoding)."""
    body = text or ""
    if "\ufffd" in body:
        return True
    if len(body) < 20:
        return False
    letters = sum(1 for c in body if c.isalpha())
    return (letters / len(body)) < GARBLE_LETTER_RATIO


def _needs_ocr(text: str) -> bool:
    if ocr_always():
        return True
    return _sparse(text) or _garbled(text)


def _ocr_pdf_page(raw: bytes, index: int) -> tuple[str, str | None]:
    try:
        import pypdfium2 as pdfium
    except Exception as exc:  # noqa: BLE001
        return "", f"ocr_skipped (pypdfium2: {exc})"
    try:
        pdf = pdfium.PdfDocument(raw)
        page = pdf[index]
        bitmap = page.render(scale=2)
        pil = bitmap.to_pil()
        pdf.close()
    except Exception as exc:  # noqa: BLE001
        return "", f"ocr_skipped (render: {exc})"
    return _ocr_pil(pil)


def _ocr_image_bytes(raw: bytes) -> tuple[str, str | None]:
    try:
        from PIL import Image

        img = Image.open(io.BytesIO(raw)).convert("RGB")
    except Exception as exc:  # noqa: BLE001
        return "", f"ocr_skipped (image: {exc})"
    return _ocr_pil(img)


def _ocr_pil(img: Any) -> tuple[str, str | None]:
    try:
        from rapidocr_onnxruntime import RapidOCR
    except Exception as exc:  # noqa: BLE001
        return "", f"ocr_skipped ({exc})"
    try:
        import numpy as np

        engine = RapidOCR()
        result, _elapse = engine(np.array(img))
    except Exception as exc:  # noqa: BLE001
        return "", f"ocr_skipped ({exc})"
    if not result:
        return "", None
    lines = []
    for row in result:
        if isinstance(row, (list, tuple)) and len(row) >= 2:
            lines.append(str(row[1]))
        elif isinstance(row, dict) and row.get("text"):
            lines.append(str(row["text"]))
    return "\n".join(lines), None


def ocr_available() -> bool:
    try:
        from rapidocr_onnxruntime import RapidOCR  # noqa: F401

        return True
    except Exception:  # noqa: BLE001
        return False
