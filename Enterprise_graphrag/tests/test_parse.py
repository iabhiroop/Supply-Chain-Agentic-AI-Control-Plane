"""PDF / OCR parse and MENTIONS ingest."""

from __future__ import annotations

import os
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from graphrag.docpack import PO_1842, QUALITY, write_scanned_pdf, write_text_pdf  # noqa: E402
from graphrag.ingest import chunk_pages  # noqa: E402
from graphrag.ontology import extract_entity_keys  # noqa: E402
from graphrag.parse import _garbled, _needs_ocr, ocr_available, parse_document  # noqa: E402


class ParseTests(unittest.TestCase):
    def test_digital_pdf_extracts_po_1842(self):
        """Feature: Digital PDF parse
        Expectation: A text PDF of PO-1842 extracts the PO and plant without OCR.
        """
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "po_1842.pdf")
            write_text_pdf(path, PO_1842.splitlines())
            parsed = parse_document(path)
        blob = parsed["text"]
        self.assertIn("PO-1842", blob)
        self.assertIn("PL-IN-CHN", extract_entity_keys(blob))
        self.assertEqual(parsed["pages"][0]["page"], 1)
        self.assertFalse(parsed["pages"][0]["ocr"])

    def test_page_anchors(self):
        chunks = chunk_pages(
            [{"page": 2, "text": "PO-1842 at PL-IN-CHN", "ocr": False}],
            heading="PO",
        )
        self.assertTrue(chunks)
        self.assertTrue(chunks[0]["anchor"].startswith("page:2#"))

    @unittest.skipUnless(ocr_available(), "rapidocr-onnxruntime not installed")
    def test_scanned_pdf_ocr_finds_vend_abc(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "quality_scan.pdf")
            write_scanned_pdf(path, QUALITY.splitlines())
            parsed = parse_document(path)
        compact = (parsed["text"] or "").upper().replace(" ", "").replace("-", "")
        self.assertIn("VENDABC", compact)
        self.assertTrue(parsed["pages"][0]["ocr"])

    def test_garbled_text_layer_needs_ocr(self):
        junk = "@@@@@@@@############$$$$$$$$%%%%"
        self.assertTrue(_garbled(junk))
        self.assertTrue(_needs_ocr(junk))
        self.assertFalse(_needs_ocr("Vendor VEND-ABC supplies MAT-SEMI-1 to plant PL-IN-CHN."))

    def test_empty_pdf_ingest_rejected(self):
        from graphrag.ingest import ingest_path

        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "empty.pdf")
            with open(path, "wb") as fh:
                fh.write(b"%PDF-1.4\n%%EOF\n")
            result = ingest_path(path)
        self.assertFalse(result.get("ok"))
        self.assertEqual(result.get("reason"), "empty_parse")
        self.assertIsNone(result.get("document_id"))

    @unittest.skipUnless(ocr_available(), "rapidocr-onnxruntime not installed")
    def test_docx_embedded_image_ocr(self):
        from PIL import Image, ImageDraw

        from graphrag.docpack import _font

        with tempfile.TemporaryDirectory() as tmp:
            png = os.path.join(tmp, "vend.png")
            img = Image.new("RGB", (900, 180), "white")
            draw = ImageDraw.Draw(img)
            draw.text((24, 50), "VEND-ABC MAT-SEMI-1", fill=(0, 0, 0), font=_font(40))
            img.save(png)
            path = os.path.join(tmp, "scan.docx")
            from docx import Document

            doc = Document()
            doc.add_paragraph("Quality scan attached.")
            doc.add_picture(png)
            doc.save(path)
            parsed = parse_document(path)
        compact = (parsed["text"] or "").upper().replace(" ", "").replace("-", "")
        self.assertIn("VENDABC", compact)
        self.assertTrue(parsed["pages"][0]["ocr"])


class MentionsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            from graphrag.apply_schema import apply
            from graphrag.db import fetch_one

            apply()
            fetch_one("SELECT 1 AS ok")
            cls.ready = True
        except Exception:  # noqa: BLE001
            cls.ready = False

    def test_ingest_pdf_mentions_po(self):
        if not self.ready:
            self.skipTest("graphrag schema / postgres unavailable")
        os.environ["GRAPHRAG_LLM_INGEST"] = "false"
        from graphrag.db import fetch_one
        from graphrag.ingest import ingest_path

        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "po_1842.pdf")
            write_text_pdf(path, PO_1842.splitlines())
            ingest_path(path)
        row = fetch_one(
            """
            SELECT n.key
              FROM edges e
              JOIN nodes n ON n.node_id = e.dst_id
             WHERE e.rel_type = 'MENTIONS' AND n.key = 'PO-1842'
             LIMIT 1
            """
        )
        self.assertIsNotNone(row)
        self.assertEqual(row["key"], "PO-1842")


if __name__ == "__main__":
    unittest.main()
